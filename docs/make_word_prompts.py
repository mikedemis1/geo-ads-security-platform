from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3)
section.right_margin  = Cm(2.5)

# ── Styles helpers ────────────────────────────────────────────────────────────
def set_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name  = name
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size=20, bold=True, color=(31, 73, 125))
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(6)

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size=12, italic=True, color=(89, 89, 89))
    p.paragraph_format.space_after = Pt(18)

def add_chapter_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=15, bold=True, color=(31, 73, 125))
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after  = Pt(6)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F497D')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=12, bold=True, color=(68, 114, 196))
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)

def add_label(doc, text):
    """Small grey label above the prompt box."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=9, italic=True, color=(128, 128, 128))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)

def add_prompt_box(doc, text):
    """Shaded box containing the prompt text."""
    p = doc.add_paragraph()
    # light blue-grey shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'EBF3FB')
    pPr.append(shd)
    # indent
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'),  '360')
    ind.set(qn('w:right'), '360')
    pPr.append(ind)

    run = p.add_run(text.strip())
    set_font(run, name="Courier New", size=9)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(8)
    return p

def add_tip(doc, text):
    p = doc.add_paragraph()
    run = p.add_run("💡 " + text)
    set_font(run, size=9, italic=True, color=(70, 130, 70))
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(4)

# ══════════════════════════════════════════════════════════════════════════════
#  COVER
# ══════════════════════════════════════════════════════════════════════════════
add_title(doc, "GEO-ADS — ChatGPT Prompts")
add_title(doc, "ανά Κεφάλαιο & Υποενότητα")
add_subtitle(doc, "Μιχάλης Δέμης, ΑΜ 1080958 | Πανεπιστήμιο Πατρών, ΤΗΜΜΥ | 2024–2025")

p = doc.add_paragraph()
run = p.add_run(
    "Οδηγίες χρήσης:\n"
    "1. Στείλε πρώτα το ΒΗΜΑ 0 (αρχικοποίηση) στο ChatGPT.\n"
    "2. Μετά στέλνεις τα prompts ένα-ένα με τη σειρά.\n"
    "3. Κάθε prompt είναι αυτοτελές — περιέχει όλο το context.\n"
    "4. Τα νούμερα στα prompts είναι ΑΚΡΙΒΗ — μην τα αλλάζεις."
)
set_font(run, size=10)
p.paragraph_format.space_after = Pt(16)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΒΗΜΑ 0
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΒΗΜΑ 0 — ΑΡΧΙΚΟΠΟΙΗΣΗ (στείλε μία φορά)")
add_tip(doc, "Στείλε αυτό πρώτο. Μία φορά ανά νέα συνεδρία ChatGPT.")
add_prompt_box(doc,
"""Είσαι βοηθός συγγραφής ακαδημαϊκής διπλωματικής εργασίας.
Θα σου στέλνω prompts ένα-ένα για κάθε υποενότητα.
Κάθε prompt περιέχει όλο το context που χρειάζεσαι.
Κανόνες:
- Γλώσσα: Ελληνικά, ακαδημαϊκό ύφος, τρίτο πρόσωπο
- Βιβλιογραφία: [1], [2]... ενδοκειμενικά, placeholder αν δεν ξέρεις DOI
- ΜΗΝ εφευρίσκεις αριθμούς — χρησιμοποίησε μόνο όσα σου δίνω
- ΜΗΝ γράφεις πλήρη κώδικα — μόνο αποσπάσματα 5-10 γραμμών
Απάντησε "Έτοιμος" για να ξεκινήσουμε.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΠΕΡΙΛΗΨΗ + ABSTRACT
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΠΕΡΙΛΗΨΗ + ABSTRACT")

add_section_heading(doc, "PROMPT P-1: Περίληψη (Ελληνικά) — 200-300 λέξεις")
add_prompt_box(doc,
"""CONTEXT:
Τίτλος: "Σύστημα Διαχείρισης Γεω-Στοχευμένων Διαφημίσεων Αθλητικών Εγκαταστάσεων σε Πραγματικό Χρόνο"
Φοιτητής: Μιχάλης Δέμης, ΑΜ 1080958, Πανεπιστήμιο Πατρών, ΤΗΜΜΥ, 2024-2025.

ΤΙ ΕΙΝΑΙ: GEO-ADS — real-time σύστημα γεω-στοχευμένης διαφήμισης σε αθλητικές εγκαταστάσεις.
Αναθέτει διαφημίσεις σε 28 ψηφιακές οθόνες γηπέδου βάσει γεωγραφικής θέσης θεατή.
Stack: FastAPI + PostgreSQL/PostGIS backend, React frontend, Electron desktop. 6 επίπεδα ασφάλειας.

ΑΠΟΤΕΛΕΣΜΑΤΑ (χρησιμοποίησε ΑΚΡΙΒΩΣ):
- R-Tree 8.09x ταχύτερο από Linear Scan για n=100.000 (1.799ms vs 14.551ms)
- Υβριδικό threat detection (rule-based + Z-score): ~0.014ms ανίχνευση

ΖΗΤΟΥΜΕΝΟ:
Γράψε Περίληψη στα Ελληνικά, 200-300 λέξεις.
Κάλυψε: (α) πρόβλημα, (β) στόχοι, (γ) μέθοδοι, (δ) αποτελέσματα, (ε) συμπέρασμα.
ΜΗΝ βάλεις βιβλιογραφία.""")

add_section_heading(doc, "PROMPT P-2: Abstract (English) — 200-300 words")
add_prompt_box(doc,
"""CONTEXT:
Title: "Real-Time Geo-Targeted Advertising Management System for Sports Facilities"
Student: Michalis Demis, University of Patras, ECE Department, 2024-2025.

System: GEO-ADS — real-time geo-targeted advertising platform for sports venues.
Assigns ads to 28 digital screens based on spectator geolocation.
Stack: FastAPI + PostgreSQL/PostGIS + React + Electron. 6 security layers.

Results:
- R-Tree index: 8.09x speedup vs Linear Scan at n=100,000 (1.799ms vs 14.551ms)
- Hybrid threat detection (rule-based + Z-score): ~0.014ms detection latency

TASK:
Write Abstract in English, 200-300 words.
Cover: (a) problem, (b) objectives, (c) methods, (d) results, (e) conclusion.
No bibliography. Academic register, passive voice preferred.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 1
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 1 — ΕΙΣΑΓΩΓΗ  (2.000–2.500 λέξεις)")

add_section_heading(doc, "PROMPT 1.1: Κίνητρο & Πρόβλημα — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS Διπλωματική:
Τίτλος: "Σύστημα Διαχείρισης Γεω-Στοχευμένων Διαφημίσεων Αθλητικών Εγκαταστάσεων σε Πραγματικό Χρόνο"

ΠΡΟΒΛΗΜΑ:
Αθλητικά γήπεδα: δεκάδες ψηφιακές οθόνες (LED, megatron, glazed floors).
Σήμερα: διαφημίσεις τυχαία ή σταθερό πρόγραμμα — χωρίς γνώση ποιος θεατής βλέπει τι.
Αποτέλεσμα: κακή στόχευση, χαμηλό ROI διαφημιστών, μη βέλτιστη χρήση χώρου.

ΛΥΣΗ (GEO-ADS):
1. Γνωρίζει γεωγραφική θέση θεατή (GPS/WiFi/BLE)
2. Αναθέτει στοχευμένες διαφημίσεις στις κοντινότερες οθόνες
3. Ενημερώνει clients ταυτόχρονα μέσω WebSocket
4. 6 επίπεδα ασφάλειας

ΖΗΤΟΥΜΕΝΟ: Γράψε υποενότητα 1.1 (~500 λέξεις).
Ξεκίνα από γενικό πρόβλημα → πρόβλημα αθλητικών εγκαταστάσεων → GEO-ADS λύση.
Βάλε [1],[2],[3] ως placeholders βιβλιογραφίας.

ΒΙΒΛΙΟΓΡΑΦΙΑ (ψάξε): "location-based advertising sports venues",
"digital signage stadiums real-time", "geo-targeted advertising ROI".""")

add_section_heading(doc, "PROMPT 1.2: Στόχοι Διπλωματικής — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 1.2 "Στόχοι Διπλωματικής"

ΣΤΟΧΟΙ:
1. Σχεδιασμός & υλοποίηση real-time γεω-στοχευμένης πλατφόρμας διαφήμισης
2. Υλοποίηση & σύγκριση 4 μεθόδων χωρικής ευρετηρίασης (Linear, R-Tree, PostGIS, Distributed)
3. Ανάπτυξη recommendation engine για αυτόματη ανάθεση διαφημίσεων σε οθόνες
4. Real-time ενημέρωση clients μέσω WebSocket protocol
5. 6 επίπεδα ασφάλειας (JWT, HMAC, Anti-Replay, Rate Limiting, Audit Log, ThreatEngine)
6. Πειραματική αξιολόγηση απόδοσης χωρικής ευρετηρίασης & ανίχνευσης απειλών
7. Συσκευασία ως all-in-one desktop εφαρμογή (Electron)

ΖΗΤΟΥΜΕΝΟ: Γράψε υποενότητα 1.2 (~300 λέξεις).
Αριθμημένη λίστα με σύντομη επεξήγηση κάθε στόχου.
Ελληνικά, ακαδημαϊκό ύφος.""")

add_section_heading(doc, "PROMPT 1.3: Ερευνητικές Ερωτήσεις ΕΡ-1 έως ΕΡ-9 — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 1.3 "Ερευνητικές Ερωτήσεις"

9 ΕΡΕΥΝΗΤΙΚΕΣ ΕΡΩΤΗΣΕΙΣ (ΑΚΡΙΒΩΣ):
ΕΡ-1: Πώς οργανώνονται ζώνες ψηφιακών οθονών σε γήπεδο; (3 ζώνες, 28 οθόνες)
ΕΡ-2: Ποια μέθοδος χωρικής ευρετηρίασης είναι πιο αποδοτική;
ΕΡ-3: Πώς υλοποιείται recommendation engine για αυτόματη ανάθεση διαφημίσεων;
ΕΡ-4: Πώς επιτυγχάνεται real-time επικοινωνία server–clients;
ΕΡ-5: Πώς προστατεύονται HTTP και WebSocket endpoints με JWT;
ΕΡ-6: Πώς ενσωματώνονται εικόνες διαφημίσεων στις ψηφιακές οθόνες;
ΕΡ-7: Γιατί PostgreSQL για real-time γεω-στοχευμένο σύστημα;
ΕΡ-8: Πώς συσκευάζεται το σύστημα ως desktop εφαρμογή με Electron;
ΕΡ-9: Πώς ανιχνεύονται απειλές με υβριδική rule-based + statistical προσέγγιση;

ΠΟΥ ΑΠΑΝΤΙΟΥΝΤΑΙ:
ΕΡ-2 → Κεφ. 6 Πείραμα 1 | ΕΡ-9 → Κεφ. 6 Πείραμα 2 | ΕΡ-3 → Κεφ. 5 | Λοιπές → Κεφ. 4+5

ΖΗΤΟΥΜΕΝΟ: Γράψε υποενότητα 1.3 (~400 λέξεις).
Παρουσίασε τις 9 ΕΡ σε πίνακα ή αριθμημένη λίστα.
Για κάθε ΕΡ: 1 πρόταση επεξήγησης + κεφάλαιο απάντησης.""")

add_section_heading(doc, "PROMPT 1.4: Συνεισφορά Εργασίας — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 1.4 "Συνεισφορά Εργασίας"

ΠΡΑΓΜΑΤΙΚΗ ΣΥΝΕΙΣΦΟΡΑ:
1. Ολοκληρωμένο real-time σύστημα γεω-στοχευμένης διαφήμισης (end-to-end)
2. Πειραματική σύγκριση 4 μεθόδων χωρικής ευρετηρίασης:
   R-Tree 8.09x ταχύτερο από Linear για n=100.000 | Distributed: ~0.44ms overhead
3. Υβριδικό threat detection (rule-based + Z-score) — πρωτότυπο για αθλητικές πλατφόρμες
4. 6-επίπεδη αρχιτεκτονική ασφάλειας
5. All-in-one Electron desktop εφαρμογή
6. Ανοιχτή, επεκτάσιμη αρχιτεκτονική (FastAPI + PostgreSQL/PostGIS)

ΖΗΤΟΥΜΕΝΟ: Γράψε υποενότητα 1.4 (~400 λέξεις).
Τι προσφέρει που δεν υπάρχει ήδη: συνδυασμός real-time + geo-targeting + security.
Βάλε [1],[2] ως placeholders. Ελληνικά, ακαδημαϊκό ύφος.""")

add_section_heading(doc, "PROMPT 1.5: Δομή Εργασίας — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 1.5 "Δομή Εργασίας"

ΚΕΦΑΛΑΙΑ:
Κεφ. 1: Εισαγωγή — κίνητρο, στόχοι, ΕΡ
Κεφ. 2: Βιβλιογραφική Ανασκόπηση — digital ads, χωρικά ευρετήρια, WebSocket, JWT, anomaly detection
Κεφ. 3: Θεωρητικό Πλαίσιο — αλγόριθμοι, WebSocket, JWT, Z-score
Κεφ. 4: Σχεδιασμός — αρχιτεκτονική, ζώνες, DB, real-time, ασφάλεια
Κεφ. 5: Υλοποίηση — backend, frontend, desktop, spatial indexing, security
Κεφ. 6: Πειράματα & Αξιολόγηση — benchmark, threat detection
Κεφ. 7: Συμπεράσματα — ΕΡ απαντήσεις, limitations, μελλοντικές προεκτάσεις

ΖΗΤΟΥΜΕΝΟ: Γράψε υποενότητα 1.5 (~300 λέξεις).
1-2 προτάσεις ανά κεφάλαιο. Τόνισε ροή: θεωρία → σχεδιασμός → υλοποίηση → αξιολόγηση.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 2
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 2 — ΒΙΒΛΙΟΓΡΑΦΙΚΗ ΑΝΑΣΚΟΠΗΣΗ  (4.500–5.500 λέξεις)")

add_section_heading(doc, "PROMPT 2.0: Εισαγωγική παράγραφος Κεφ. 2 — ~150 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS Βιβλιογραφική Ανασκόπηση

ΖΗΤΟΥΜΕΝΟ: Γράψε εισαγωγική παράγραφο Κεφ. 2 (~150 λέξεις).
Εξήγησε ότι καλύπτει 5 θεματικές:
(α) ψηφιακή διαφήμιση & γεω-στόχευση, (β) χωρικά ευρετήρια,
(γ) real-time αρχιτεκτονικές & WebSocket, (δ) ασφάλεια web,
(ε) ανίχνευση ανωμαλιών.
Κάθε θεματική συνδέεται με ΕΡ-1 έως ΕΡ-9.""")

add_section_heading(doc, "PROMPT 2.1: Ψηφιακή Διαφήμιση & Γεω-Στόχευση — ~1.000 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 2.1

ΣΧΕΣΗ ΜΕ ΕΡΓΟ: 28 οθόνες 3 ζωνών, ανάθεση βάσει θέσης θεατή.

ΘΕΜΑΤΑ:
1. Εξέλιξη digital advertising: static → programmatic/real-time
2. Location-based advertising (LBA): geofencing, proximity
3. Digital signage σε αθλητικές εγκαταστάσεις
4. Out-of-Home (OOH) → Programmatic Digital OOH (pDOOH)
5. Σύγκριση, κενό γνώσης

ΖΗΤΟΥΜΕΝΟ: ~1.000 λέξεις, θεματική οργάνωση, σύγκριση μεθόδων.
Κλείσε με "Κενό Γνώσης" — τι λείπει (που καλύπτει το GEO-ADS).
Βιβλιογραφία [1]-[6] placeholders.

ΠΗΓΕΣ (ψάξε): "programmatic digital out-of-home advertising" (2019-2024),
"location-based advertising mobile", "digital signage sports stadiums",
"geo-targeted advertising effectiveness ROI".""")

add_section_heading(doc, "PROMPT 2.2: Χωρικά Ευρετήρια — ~1.000 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 2.2

BENCHMARK (ΑΚΡΙΒΩΣ):
n=28: Linear 0.004ms, R-Tree 0.008ms | n=100.000: Linear 14.551ms, R-Tree 1.799ms

ΘΕΜΑΤΑ:
1. Ιστορία: B-Tree → R-Tree (Guttman 1984)
2. R-Tree παραλλαγές: R*-Tree, R+-Tree, Hilbert R-Tree
3. PostGIS και GiST ευρετήρια
4. Κατανεμημένη χωρική επεξεργασία
5. Σύγκριση, κενό γνώσης

ΖΗΤΟΥΜΕΝΟ: ~1.000 λέξεις, εξέλιξη μεθόδων, σύγκριση, κενό γνώσης.
Βιβλιογραφία [1]-[7] placeholders.

ΠΗΓΕΣ: Guttman 1984 R-Trees SIGMOD, Beckmann 1990 R*-Tree SIGMOD,
"spatial database indexing comparison" surveys, PostGIS docs.""")

add_section_heading(doc, "PROMPT 2.3: Real-time Αρχιτεκτονικές & WebSocket — ~1.000 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 2.3

ΣΤΟΙΧΕΙΑ ΕΡΓΟΥ: 3 WS channels (/ws/placements, /ws/layout, /ws/security).
Backend: FastAPI asyncio. Frontend: React native WebSocket API, auto-reconnect.

ΘΕΜΑΤΑ:
1. Εξέλιξη: polling → long-polling → SSE → WebSocket
2. WebSocket RFC 6455: handshake, frames, full-duplex
3. WebSocket vs REST για real-time
4. Event-driven αρχιτεκτονικές: pub-sub
5. ASGI/FastAPI για async WebSocket

ΖΗΤΟΥΜΕΝΟ: ~1.000 λέξεις, σύγκριση τεχνολογιών, κενό γνώσης.
Βιβλιογραφία [1]-[6].

ΠΗΓΕΣ: RFC 6455 IETF 2011, "WebSocket performance comparison polling SSE",
"real-time web applications architecture" surveys.""")

add_section_heading(doc, "PROMPT 2.4: Ασφάλεια Web Εφαρμογών — ~1.000 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 2.4

6 ΕΠΙΠΕΔΑ ΑΣΦΑΛΕΙΑΣ:
v1: JWT HS256, 1h access + 24h refresh, 6 scopes
v2: HMAC-SHA256 WS μηνύματα
v3: Anti-Replay (60s + nonce dedup)
v4: Rate Limiting (10/min login, 5/min token)
v5: Audit Log JSON
v6: ThreatEngine (rule + Z-score)

ΘΕΜΑΤΑ:
1. JWT RFC 7519: δομή, claims, HS256 vs RS256, scopes
2. HMAC RFC 2104: HMAC(K,m) φόρμουλα
3. Replay attacks + αντίμετρα
4. Rate limiting: Token Bucket vs Sliding Window
5. OWASP Top 10 — σχετικά με το εργο
6. WebSocket security challenges

ΖΗΤΟΥΜΕΝΟ: ~1.000 λέξεις, σύγκριση μηχανισμών, κενό γνώσης.

ΠΗΓΕΣ: RFC 7519, RFC 2104, OWASP Top Ten 2021, "JWT security vulnerabilities" surveys.""")

add_section_heading(doc, "PROMPT 2.5: Ανίχνευση Ανωμαλιών — ~1.000 λέξεις")
add_prompt_box(doc,
"""CONTEXT — GEO-ADS, υποενότητα 2.5

THREAT ENGINE:
Rule-based: brute_force (5/5min/IP→CRITICAL), credential_stuffing (10 usernames/10min),
            replay_attack (3/5min), rate_abuse (20/10min→WARNING)
Statistical: z=(rate-mean)/std, alert z>2.0, windows 5m/15m/1h
Αποτελέσματα: Rule ~0.014ms, Statistical ~0.013ms

ΘΕΜΑΤΑ:
1. IDS taxonomy: signature-based vs anomaly-based
2. Rule-based: πλεονεκτήματα/μειονεκτήματα
3. Statistical Z-score, baseline, sliding windows
4. ML για threat detection vs rule-based
5. Υβριδικές προσεγγίσεις

ΖΗΤΟΥΜΕΝΟ: ~1.000 λέξεις, σύγκριση, γιατί ο υβριδισμός βέλτιστος.

ΠΗΓΕΣ: "anomaly detection intrusion detection survey" (2020-2024),
"Z-score statistical anomaly detection web", "brute force detection sliding window".""")

add_section_heading(doc, "PROMPT 2.6: Σύνοψη Κεφ. 2 — ~200 λέξεις")
add_prompt_box(doc,
"""ΖΗΤΟΥΜΕΝΟ: Γράψε σύνοψη ~200 λέξεις για το τέλος Κεφ. 2.
Συνόψισε κυριότερα ευρήματα 2.1-2.5.
Τόνισε "κενό γνώσης" που καλύπτει το GEO-ADS.
Μετάβαση στο Κεφ. 3 (Θεωρητικό Πλαίσιο).""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 3
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 3 — ΘΕΩΡΗΤΙΚΟ ΠΛΑΙΣΙΟ  (5.000–6.000 λέξεις)")

add_section_heading(doc, "PROMPT 3.1.1: Linear Scan O(n) — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.1.1 "Linear Scan"
Loop όλων των n screens, Haversine distance, filter εντός radius. O(n).
BENCHMARK: n=28: 0.004ms | n=100: 0.018ms | n=10K: 1.475ms | n=100K: 14.551ms

ΖΗΤΟΥΜΕΝΟ: ~300 λέξεις. Αλγόριθμος, πολυπλοκότητα, Haversine, πότε κατάλληλος.""")

add_section_heading(doc, "PROMPT 3.1.2: R-Tree O(log n + k) — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.1.2 "R-Tree"
In-memory R-Tree (Python rtree/libspatialindex). MBR για κάθε screen.
Insertion O(logn), Range Query O(logn+k). Guttman 1984.
BENCHMARK: n=28: 0.008ms | n=100: 0.012ms | n=10K: 0.200ms | n=100K: 1.799ms (8.09x speedup)

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις.
(α) Δομή R-Tree + MBR, (β) αλγόριθμοι insertion/query, (γ) πολυπλοκότητα,
(δ) γιατί υπερτερεί Linear για μεγάλα n.
Βιβλιογραφία: [1] Guttman 1984, [2] Beckmann 1990 R*-Tree.""")

add_section_heading(doc, "PROMPT 3.1.3: PostGIS GIST — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.1.3 "PostGIS GIST"
GEOGRAPHY(Point,4326) column. CREATE INDEX ... USING GIST(location).
Query: ST_DWithin(location, ST_MakePoint(lon,lat)::geography, radius_meters)
GiST = Generalized Search Tree. WGS-84 (EPSG:4326). Disk I/O overhead.

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις.
(α) GiST index, (β) PostGIS ST_* functions, (γ) WGS-84, (δ) disk I/O vs in-memory.""")

add_section_heading(doc, "PROMPT 3.1.4: Distributed — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.1.4 "Distributed Spatial Index"
Python ThreadPoolExecutor, 3 shards (⅓ screens each), parallel R-Tree, merge results.
MapReduce pattern. Σταθερό overhead ~0.44ms από thread spawning.
n=28: 0.444ms (111x αργότερο!), n=100K: 2.414ms (χειρότερο από single R-Tree).

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις. MapReduce, sharding, Python GIL overhead,
πότε αξίζει distributed (πραγματικά distributed systems).""")

add_section_heading(doc, "PROMPT 3.1.5: Σύγκριση Μεθόδων — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.1.5 "Σύγκριση"
| Μέθοδος     | Complexity    | Κατάλληλο  |
| Linear      | O(n)          | n < 100    |
| R-Tree      | O(log n + k)  | n > 100    |
| PostGIS GIST| O(log n)+I/O  | Persistent |
| Distributed | O(log n/k)par | n>10M multi|

n=100K: Linear 14.551ms, R-Tree 1.799ms, Distributed 2.414ms
Επιλογή production: R-Tree (O(log n), in-memory, επεκτάσιμο).

ΖΗΤΟΥΜΕΝΟ: ~300 λέξεις. Πίνακας σύγκρισης + αιτιολόγηση επιλογής R-Tree.""")

add_section_heading(doc, "PROMPT 3.2: WebSocket Protocol — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.2
3 WS channels: /ws/placements, /ws/layout, /ws/security (JWT ως query param).
HMAC-SHA256 υπογραφή μηνυμάτων. RFC 6455.

ΘΕΩΡΙΑ: Handshake (HTTP Upgrade), frames (FIN/opcode/mask/payload),
full-duplex, ping/pong, vs polling/SSE.

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις.
(α) handshake, (β) frame structure, (γ) auth challenges, (δ) WebSocket vs REST.""")

add_section_heading(doc, "PROMPT 3.3: JWT & HMAC — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.3
JWT HS256: access 1h, refresh 24h.
Scopes: ads:read, placements:read, placements:write, layout:read, recommendation:read, security:read.
JWT_SECRET >= 32 bytes (crash αν λιγότερο).
HMAC: HMAC(K,m) = H((K⊕opad)||H((K⊕ipad)||m))
Anti-Replay: 60s timestamp window + nonce dedup.

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις.
(α) JWT δομή+claims, (β) HS256, (γ) HMAC φόρμουλα, (δ) Anti-Replay, (ε) scopes.""")

add_section_heading(doc, "PROMPT 3.4: Z-score & Ανίχνευση Απειλών — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 3.4
z = (rate - mean) / std, alert αν z > 2.0, windows 5min/15min/1h, warm-up >=3 points.
Rule-based: 4 κανόνες (brute_force/credential_stuffing/replay_attack/rate_abuse).
Υβριδισμός: τρέχουν ΠΑΡΑΛΛΗΛΑ.

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις.
(α) Z-score ορισμός, (β) sliding windows, (γ) threshold z=2.0,
(δ) rule vs statistical πλεονεκτήματα, (ε) σκεπτικό υβριδισμού.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 4
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 4 — ΣΧΕΔΙΑΣΜΟΣ ΣΥΣΤΗΜΑΤΟΣ  (3.000–4.000 λέξεις)")

add_section_heading(doc, "PROMPT 4.1: Αρχιτεκτονική Συστήματος — ~600 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 4.1
4-επίπεδη αρχιτεκτονική:
Presentation: Electron → React 18 (port 3000) → VisualBoard + SecurityDashboard
Application: FastAPI (port 8000) → REST + WebSocket
Data: PostgreSQL 16 + PostGIS (port 5433, Docker)
Security: JWT + HMAC + Anti-Replay + RateLimit + AuditLog + ThreatEngine

Ροή: Θεατής check-in → FastAPI → R-Tree → Recommendation → WS broadcast → React

ΖΗΤΟΥΜΕΝΟ: ~600 λέξεις. 4 επίπεδα, ροές δεδομένων, protocols.
Αναφέρσου σε Σχήμα 4.1 (θα φτιαχτεί). Αιτιολόγησε επιλογές τεχνολογίας.""")

add_section_heading(doc, "PROMPT 4.2: Ζώνες Γηπέδου & Μοντέλο Οθονών — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 4.2
ΖΩΝΕΣ (ΑΚΡΙΒΩΣ):
GlassFloor: 4×4 = 16 tiles, glassfloor_tile, γυάλινο δάπεδο
Surrounding: 2×4 = 8 banners, surrounding_banner, περιμετρικές
Megatron: 2×2 = 4 panels, megatron_panel, κεντρική μεγαοθόνη

DB: screens(id, zone_id FK→zones, row_idx, col_idx, screen_type, location GEOGRAPHY(Point,4326))
CREATE INDEX idx_screens_location ON screens USING GIST(location);

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις. Περιγραφή ζωνών, data model, GEOGRAPHY vs GEOMETRY,
Σχήμα 4.2 (διάγραμμα γηπέδου).""")

add_section_heading(doc, "PROMPT 4.3: Σχεδιασμός Βάσης Δεδομένων — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 4.3
SCHEMA:
advertisements(id SERIAL PK, name VARCHAR, image_url TEXT, zone VARCHAR, created_at TIMESTAMP)
zones(id SERIAL PK, name VARCHAR, type VARCHAR)
screens(id SERIAL PK, zone_id FK, row_idx INT, col_idx INT, screen_type VARCHAR, location GEOGRAPHY)
Placements: in-memory (δεν αποθηκεύονται — γνωστό limitation)
28 screens pre-seeded με αναλογικές GPS συντεταγμένες.

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις. ER diagram description (Σχήμα 4.3), πίνακες, σχέσεις,
αιτιολόγηση PostgreSQL vs NoSQL, in-memory limitation.""")

add_section_heading(doc, "PROMPT 4.4: Σχεδιασμός Real-time Επικοινωνίας — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 4.4
3 CHANNELS:
/ws/placements: payload {screen_id, ad_id, image_url, zone, timestamp} → VisualBoard
/ws/layout: layout changes
/ws/security (scope:security:read): {alert_id, type, severity, message, timestamp} → SecurityDashboard
Auth: JWT ως query parameter (WebSocket δεν υποστηρίζει Authorization header).

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις. 3 channels, payloads, lifecycle, γιατί WS vs polling,
JWT-in-query-param και security implications.""")

add_section_heading(doc, "PROMPT 4.5: Σχεδιασμός Ασφάλειας — 6 Επίπεδα — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 4.5
v1 JWT: HS256, 1h access / 24h refresh, 6 scopes
v2 HMAC-SHA256: υπογραφή WS μηνυμάτων
v3 Anti-Replay: 60s timestamp + nonce dedup
v4 Rate Limiting: 10/min login, 5/min /auth/token, 10/min refresh
v5 Audit Log: JSON middleware → audit.log
v6 Threat Engine: rule(4 κανόνες) + Z-score statistical, WS push

AUTH FLOW: POST /auth/login → {access_token(1h), refresh_token(24h)}
           POST /auth/refresh → access_token(1h)
           POST /auth/token (X-Admin-Secret) → JWT (internal)

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις. Defense-in-depth, κάθε επίπεδο: σκοπός + ποιες επιθέσεις.
Σχήμα 4.5 (auth flow). OWASP σύνδεση.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 5
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 5 — ΥΛΟΠΟΙΗΣΗ ΣΥΣΤΗΜΑΤΟΣ  (4.500–5.500 λέξεις)")

add_section_heading(doc, "PROMPT 5.1: Εργαλεία & Αιτιολόγηση — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.1
Backend: FastAPI(Python 3.11+), psycopg2, python-jose, slowapi, rtree/libspatialindex, uvicorn
DB: PostgreSQL 16 + PostGIS, Docker
Frontend: React 18, native WebSocket API
Desktop: Electron, cross-env
Security: JWT HS256, HMAC-SHA256, hmac stdlib, threading.Lock

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις. Πίνακας: [Τεχνολογία | Σκοπός | Εναλλακτική | Λόγος Επιλογής].
FastAPI: async, OpenAPI auto, WS native. React: state+hooks real-time friendly.
PostgreSQL: ACID+PostGIS. Electron: all-in-one executable.""")

add_section_heading(doc, "PROMPT 5.2: Backend — REST API & WebSocket — ~600 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.2
ENDPOINTS: GET /ads (ads:read), GET /layout (layout:read), POST /placements (placements:write),
GET /recommendations?lat=&lon=&radius= (recommendation:read),
GET /security/alerts,/stats,/comparison (security:read),
POST /auth/login, /auth/refresh, /auth/token
WS: /ws/placements, /ws/layout, /ws/security

KOD ΑΠΟΣΠΑΣΜΑ (βάλε verbatim):
@app.get("/recommendations")
async def get_recommendations(
    lat: float = Query(..., ge=-90, le=90),
    lon: float = Query(..., ge=-180, le=180),
    radius: float = Query(100.0, ge=1, le=50000),
    current_user: dict = Depends(require_scope("recommendation:read"))
):
    results = layout_service.find_screens_rtree(lat, lon, radius)
    return {"screens": results, "count": len(results)}

ΖΗΤΟΥΜΕΝΟ: ~600 λέξεις. Endpoints, input validation (Pydantic ge/le),
dependency injection, scopes (principle of least privilege).""")

add_section_heading(doc, "PROMPT 5.3: Spatial Indexing — 4 Μέθοδοι — ~600 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.3 (layout_service.py)

1. find_screens_linear: loop+Haversine, O(n)
2. find_screens_rtree: bbox query→Haversine validation, O(log n+k)
3. find_screens_postgis: ST_DWithin(%s,%s,%s) parameterized (fix SQL injection!), GIST auto
4. find_screens_distributed: ThreadPoolExecutor 3 shards, merge+dedup

KOD ΑΠΟΣΠΑΣΜΑ (βάλε verbatim):
# R-Tree query
def find_screens_rtree(self, lat, lon, radius_m):
    bbox = self._expand_bbox(lat, lon, radius_m)
    candidates = list(self.rtree_idx.intersection(bbox, objects=True))
    return [s for s in candidates
            if haversine(lat, lon, s.object['lat'], s.object['lon']) <= radius_m]

ΖΗΤΟΥΜΕΝΟ: ~600 λέξεις. Κάθε μέθοδος: υλοποίηση, trade-offs.
Σχολίασε parameterized query (SQL injection prevention).""")

add_section_heading(doc, "PROMPT 5.4: Recommendation Engine — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.4
ΛΟΓΙΚΗ: θέση θεατή + radius → R-Tree → screens ταξινομημένα κατά απόσταση →
διαθέσιμα ads ανά zone → WS broadcast σε όλους.
In-memory: {screen_id → ad_id}. Zone affinity για ads.
ΠΕΡΙΟΡΙΣΜΟΙ: χωρίς ML, χωρίς personalization, in-memory.

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις. Αλγόριθμος, ροή query→broadcast,
περιορισμοί, μελλοντικό ML recommendation.""")

add_section_heading(doc, "PROMPT 5.5: Frontend — VisualBoard & SecurityDashboard — ~500 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.5
VISUALBOARD: 3 grids (GlassFloor 4×4, Surrounding 2×4, Megatron 2×2),
χρωματισμός tiles κατά assigned ad, WS /ws/placements, click→ad details,
auto-reconnect, JWT auto-refresh 2min πριν λήξη.

SECURITYDASHBOARD: 4 panels:
1. Live Alerts (severity badge CRITICAL/WARNING/INFO)
2. Event Timeline
3. Method Comparison (rule vs statistical)
4. Event Log
WS /ws/security, πολύχρωμα severity indicators.

TAB Nav App.js: Tab1 Stadium→VisualBoard, Tab2 Security→SecurityDashboard.

ΖΗΤΟΥΜΕΝΟ: ~500 λέξεις. Κάθε component, real-time behavior.
Αναφέρσου σε screenshots Σχήμα 5.5a/5.5b. Token auto-refresh mechanism.""")

add_section_heading(doc, "PROMPT 5.6: Desktop Εφαρμογή Electron — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.6
Electron wrapper: εκκινεί FastAPI subprocess + React, native window, process cleanup.
npm run desktop. Cross-platform.

ΖΗΤΟΥΜΕΝΟ: ~300 λέξεις. Σκοπός, τι επιλύει (dependency management, easy deploy).
Σύγκριση: web-only vs Docker vs Electron.""")

add_section_heading(doc, "PROMPT 5.7: Σύστημα Ασφάλειας — Υλοποίηση — ~700 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 5.7
ΑΡΧΕΙΑ: auth_config.py (crash αν JWT_SECRET<32b), auth_routes.py, jwt_service.py,
deps.py (require_scope), threat_engine.py (ThreatEngine singleton), message_schema.py

KOD (βάλε verbatim):
def require_scope(scope: str):
    def dependency(token: str = Depends(oauth2_scheme)):
        payload = verify_token(token)
        if scope not in payload.get("scopes", []):
            raise HTTPException(403, "Insufficient scope")
        return payload
    return dependency

RATE LIMITING: @limiter.limit("10/minute") στο /auth/login → 429 αν υπερβεί
THREAT ENGINE: ThreatEngine singleton → process_event(type,ip,username) →
  rule check + z-score → ThreatAlert list → WS push to /ws/security
SECURITY HEADERS: X-Frame-Options:DENY, X-Content-Type-Options:nosniff,
  X-XSS-Protection:1;mode=block, CSP:default-src 'self'

ΖΗΤΟΥΜΕΝΟ: ~700 λέξεις. Κάθε επίπεδο με κώδικα αποσπάσματα,
security headers σημασία, ThreatEngine singleton pattern.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 6
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 6 — ΠΕΙΡΑΜΑΤΑ & ΑΞΙΟΛΟΓΗΣΗ  (4.000–5.000 λέξεις)")

add_section_heading(doc, "PROMPT 6.1: Μεθοδολογία Πειραμάτων — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 6.1
Hardware: Windows 10 Pro, τοπικό μηχάνημα | Software: Python 3.11+, PostgreSQL 16, React 18
Benchmark: benchmark_all.py | Επαναλήψεις: 1.000/N (100 warm-up απορρίπτονται)
Metrics: μέση τιμή (ms), τυπική απόκλιση, speedup ratio
N = {28, 100, 1.000, 10.000, 100.000}, ακτίνα query: 500m σταθερή

Threat test: test_threat_detection.py, 36 tests, 8 κατηγορίες:
Auth, JWT, Rate limiting, Replay, Brute force, Credential stuffing, Security API, WS security

ΖΗΤΟΥΜΕΝΟ: ~300 λέξεις. Hardware/software/params, δικαιολόγηση 1.000 επαναλήψεων,
περιορισμοί (τοπικό μηχάνημα, mock GPS).""")

add_section_heading(doc, "PROMPT 6.2: Πείραμα 1 — Χωρική Ευρετηρίαση — ~800 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 6.2 (ΕΡ-2)
ΑΠΟΤΕΛΕΣΜΑΤΑ (ΑΚΡΙΒΩΣ):
| N       | Linear (ms) | R-Tree (ms) | Distributed (ms) | Speedup R-Tree |
| 28      | 0.004       | 0.008       | 0.444            | 0.54x          |
| 100     | 0.018       | 0.012       | 0.435            | 1.45x          |
| 1.000   | 0.148       | 0.035       | 0.455            | 4.25x          |
| 10.000  | 1.475       | 0.200       | 0.616            | 7.36x          |
| 100.000 | 14.551      | 1.799       | 2.414            | 8.09x          |

ΕΡΜΗΝΕΙΑ:
n=28: Linear κερδίζει (overhead R-Tree > benefit), crossover ~n=50-80
n=100K: R-Tree 8.09x ταχύτερο. Θεωρία: O(n)→100K βήματα vs O(logn)→17 βήματα
Distributed: thread spawning Python GIL = ~0.44ms overhead, n=28: 111x αργότερο!

ΖΗΤΟΥΜΕΝΟ: ~800 λέξεις σε 4 sub-sections:
6.2.1 Στόχος (~100λ) | 6.2.2 Μεθοδολογία (~150λ) | 6.2.3 Αποτελέσματα+πίνακας (~200λ)
6.2.4 Ανάλυση O(n) vs O(logn), crossover, distributed (~350λ)
Αναφέρσου σε Σχήμα 6.1 (χρόνος vs N) + Σχήμα 6.2 (speedup). Απάντησε στην ΕΡ-2.""")

add_section_heading(doc, "PROMPT 6.3: Πείραμα 2 — Threat Detection — ~700 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 6.3 (ΕΡ-9)
ΑΠΟΤΕΛΕΣΜΑΤΑ:
Rule avg: ~0.014ms | Statistical avg: ~0.013ms
Rule TP: brute_force, credential_stuffing, replay (known patterns) = 100% TP
Statistical TP: gradual anomaly, low-and-slow attacks = ~85% TP, ~15% FP
Rule miss: novel attacks, gradual change
Statistical miss: sudden burst (κάτω από threshold windows)

ΚΑΝΟΝΕΣ: brute_force(5/5min/IP→CRITICAL), credential_stuffing(10usernames/10min),
         replay_attack(3/5min), rate_abuse(20/10min→WARNING)
Z-score: z=(rate-mean)/std > 2.0, windows 5m/15m/1h, warm-up >=3

ΖΗΤΟΥΜΕΝΟ: ~700 λέξεις σε 4 sub-sections:
6.3.1 Στόχος | 6.3.2 Μεθοδολογία | 6.3.3 Αποτελέσματα (TP/FP πίνακας, χρόνοι)
6.3.4 Ανάλυση (πότε κερδίζει κάθε μέθοδος, συμπληρωματικότητα)
Απάντησε στην ΕΡ-9.""")

add_section_heading(doc, "PROMPT 6.4: Συζήτηση Αποτελεσμάτων — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 6.4
Πείραμα 1: R-Tree βέλτιστο production (8.09x, O(logn)). Linear αρκετό για n=28.
           Distributed: single-machine overhead κυριαρχεί.
Πείραμα 2: Υβριδισμός υπερέχει. Rule: deterministic known. Statistical: novel/gradual.

ΠΕΡΙΟΡΙΣΜΟΙ: Windows τοπικό, n>100 προσομοίωση, ελεγχόμενα threat scenarios, warm-up 3pts.

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις. Συνδυασμός ευρημάτων, περιορισμοί, validity, σχέση με Κεφ. 2.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΚΕΦΑΛΑΙΟ 7
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΚΕΦΑΛΑΙΟ 7 — ΣΥΜΠΕΡΑΣΜΑΤΑ & ΜΕΛΛΟΝΤΙΚΕΣ ΠΡΟΕΚΤΑΣΕΙΣ  (1.500–2.000 λέξεις)")

add_section_heading(doc, "PROMPT 7.1: Ανακεφαλαίωση — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 7.1
Υλοποιήθηκε: GEO-ADS end-to-end (3 ζώνες, 28 οθόνες, FastAPI+PostgreSQL+React+Electron).
4 μέθοδοι χωρικής ευρετηρίασης (benchmark). 6 επίπεδα ασφάλειας.
Υβριδικό threat detection. WebSocket real-time.

ΖΗΤΟΥΜΕΝΟ: ~300 λέξεις. Σύντομη ανακεφαλαίωση συνδέοντας με στόχους Κεφ. 1.""")

add_section_heading(doc, "PROMPT 7.2: Απαντήσεις στις ΕΡ-1 έως ΕΡ-9 — ~600 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 7.2
ΕΡ-1: 3 ζώνες (GlassFloor 4×4, Surrounding 2×4, Megatron 2×2), screen_type + GPS.
ΕΡ-2: R-Tree 8.09x speedup για n=100K (1.799ms vs 14.551ms). Για n=28 Linear ταχύτερο.
ΕΡ-3: Proximity-based: R-Tree → ταξινόμηση απόσταση → zone affinity → WS broadcast.
ΕΡ-4: WebSocket full-duplex (RFC 6455), 3 channels, <1ms latency.
ΕΡ-5: JWT Bearer (HS256, 6 scopes) HTTP. WS: JWT query param + HMAC-SHA256.
ΕΡ-6: advertisements.image_url στη DB, frontend φορτώνει σε tile κατά assignment.
ΕΡ-7: ACID + PostGIS (ST_DWithin, GIST) + ωριμότητα production.
ΕΡ-8: Electron εκκινεί FastAPI+React subprocess, all-in-one χωρίς εξαρτήσεις.
ΕΡ-9: Rule(4 κανόνες, ~0.014ms) + Z-score(z>2.0, 5m/15m/1h, ~0.013ms) παράλληλα.

ΖΗΤΟΥΜΕΝΟ: ~600 λέξεις. Κάθε ΕΡ σαφής απάντηση με πειραματικά για ΕΡ-2 και ΕΡ-9.""")

add_section_heading(doc, "PROMPT 7.3: Limitations — ~300 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 7.3
ΓΝΩΣΤΟΙ ΠΕΡΙΟΡΙΣΜΟΙ:
1. Placements in-memory (χάνονται στο restart)
2. Benchmark τοπικό Windows — μη αντιπροσωπευτικό production
3. n>100 screens: τεχνητή προσομοίωση (πραγματικό: n=28)
4. Statistical warm-up >= 3 data points
5. Χωρίς real GPS (mock συντεταγμένες)
6. Χωρίς personalization (μόνο proximity)
7. Single-instance (χωρίς load balancing)

ΖΗΤΟΥΜΕΝΟ: ~300 λέξεις. Ειλικρινής παρουσίαση, γιατί δεν αναιρούν τα ευρήματα.""")

add_section_heading(doc, "PROMPT 7.4: Μελλοντικές Προεκτάσεις — ~400 λέξεις")
add_prompt_box(doc,
"""CONTEXT — 7.4
ΕΠΕΚΤΑΣΕΙΣ:
Βραχυπρόθεσμες: Persistence placements στη DB, ML Recommendation (collaborative filtering/RL)
Μεσοπρόθεσμες: ML Threat Detection (LSTM/Isolation Forest), Multi-venue multi-tenant
Μακροπρόθεσμες: Mobile app (React Native), Real GPS (BLE/WiFi/UWB),
               Horizontal scaling (Redis pub-sub), A/B Testing, Analytics dashboard (CTR/heatmaps)

ΖΗΤΟΥΜΕΝΟ: ~400 λέξεις. Ομαδοποίηση βραχυ/μακροπρόθεσμες, αιτιολόγηση σημασίας.""")

doc.add_page_break()

# ══════════════════════════════════════════════════════════════════════════════
#  ΒΙΒΛΙΟΓΡΑΦΙΑ + EXTRAS
# ══════════════════════════════════════════════════════════════════════════════
add_chapter_heading(doc, "ΒΙΒΛΙΟΓΡΑΦΙΑ — Αιτήματα ανά Κεφάλαιο")

add_section_heading(doc, "PROMPT BIB-2: Βιβλιογραφία Κεφ. 2 (20-25 πηγές)")
add_prompt_box(doc,
"""Ψάξε 20-25 επιστημονικές αναφορές για τη βιβλιογραφία Κεφ. 2 (GEO-ADS):

1. Digital advertising & geo-targeting (5 πηγές, 2019-2024)
2. Spatial indexing & R-Tree (5 πηγές, classic + modern)
3. WebSocket & real-time architectures (4 πηγές)
4. Web security: JWT, HMAC, OWASP (4 πηγές)
5. Anomaly detection & IDS (5 πηγές, 2020-2024)

Μορφή: [N] Συγγραφέας. "Τίτλος." Περιοδικό, Χρονολογία. doi: https://doi.org/...
Σημείωσε unavailable DOI (RFC, books).""")

add_section_heading(doc, "PROMPT BIB-3: Βιβλιογραφία Κεφ. 3 (15 πηγές)")
add_prompt_box(doc,
"""Ψάξε 15 πηγές για Κεφ. 3 (Θεωρητικό Πλαίσιο):
1. Guttman R-Tree 1984 + R*-Tree + B-Tree + spatial complexity
2. WebSocket RFC 6455 + ASGI
3. JWT RFC 7519 + HMAC RFC 2104
4. Z-score anomaly detection, sliding windows

Μορφή: [N] Συγγραφέας. "Τίτλος." Περιοδικό, Χρονολογία. doi: ...""")

add_section_heading(doc, "PROMPT BIB-6: Βιβλιογραφία Κεφ. 6 (8-10 πηγές)")
add_prompt_box(doc,
"""Ψάξε 8-10 πηγές για Κεφ. 6 (Πειράματα):
1. Benchmark methodology για spatial databases
2. R-Tree performance evaluation
3. IDS evaluation metrics (TP/FP/FN)
4. Statistical anomaly detection benchmarking

Μορφή: [N] Συγγραφέας. "Τίτλος." Περιοδικό, Χρονολογία. doi: ...""")

add_chapter_heading(doc, "EXTRA PROMPTS")

add_section_heading(doc, "PROMPT E-1: Πίνακας Περιεχομένων")
add_prompt_box(doc,
"""Δημιούργησε Πίνακα Περιεχομένων για τη διπλωματική GEO-ADS (100-120 σελίδες).
Κεφ. 1 Εισαγωγή (~10σ), Κεφ. 2 Βιβλιογραφία (~25σ), Κεφ. 3 Θεωρητικό (~28σ),
Κεφ. 4 Σχεδιασμός (~18σ), Κεφ. 5 Υλοποίηση (~25σ), Κεφ. 6 Πειράματα (~22σ), Κεφ. 7 Συμπεράσματα (~10σ).
Μορφή: Αριθμός | Τίτλος | Σελίδα. Συμπεριέλαβε όλες τις υποενότητες.""")

add_section_heading(doc, "PROMPT E-2: Ευρετήριο Εικόνων & Πινάκων")
add_prompt_box(doc,
"""Δημιούργησε λίστα Σχημάτων & Πινάκων GEO-ADS με λεζάντες (15-20 λέξεις, ελληνικά):
Σχ. 1.1: Architecture overview | Σχ. 3.1: R-Tree MBR | Σχ. 3.2: WebSocket handshake
Σχ. 4.1: System architecture | Σχ. 4.2: Stadium zones | Σχ. 4.3: ER diagram
Σχ. 4.5: Auth flow | Σχ. 5.5a: VisualBoard screenshot | Σχ. 5.5b: SecurityDashboard screenshot
Σχ. 6.1: Time vs N graph | Σχ. 6.2: Speedup vs N graph
Πίν. 3.1: Σύγκριση πολυπλοκότητας | Πίν. 5.1: Τεχνολογίες | Πίν. 6.1: Benchmark ms | Πίν. 6.2: Threat detection""")

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = r"C:\Projects\MIKE_DIPLOMA\docs\GEO-ADS_ChatGPT_Prompts.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
